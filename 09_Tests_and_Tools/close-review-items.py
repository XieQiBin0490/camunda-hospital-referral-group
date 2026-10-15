#!/usr/bin/env python3
"""Close out the review list in the product backlog.

Items 1, 2 and 3 of the review list are delivered, so their backlog rows stop being
"Not started" / "In progress" and carry the evidence that closes them.

usage: python close-review-items.py [--check]
"""

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "01_Agile_Planning" / (
    "Hospital_Patient_System_Project_Backlog_按评价逐项修正最终版.xlsx")

DONE = "Done\n完成"

# id -> the evidence that closes it
PROJECT = {
    "PB-45": ("Persistence delivered and verified: support/Database.java plus H2, "
              "HospitalStore and OperationsLog re-backed by tables with their public API "
              "unchanged. PersistenceCheck writes in one JVM and reads back in a second: "
              "7 of 7 checks pass, and 41 of 41 unit tests still pass. Evidence: "
              "05_Test_Evidence/persistence/, 08_Repository_Documents/database-integration.md"),
    "PB-46": ("RBAC in force: the cluster had authorizations disabled and API protection "
              "skipped, so every authenticated user saw every task. Both settings corrected; "
              "10 candidate groups and 10 role users created from the models, with group "
              "authorizations. verify-rbac.py: membership 10/10, an ungranted user sees 0 "
              "tasks where a granted role user sees 785, and each role's inbox routes only "
              "its own tasks. Evidence: 05_Test_Evidence/rbac/, "
              "08_Repository_Documents/rbac-setup.md"),
    "PB-47": ("All four open model defects fixed and deployed: the refund decision now has "
              "an exclusive gateway, the urgent-treatment AND-split is explicit, S1 branches "
              "on referralDecision from its own form, and the misplaced and missing form "
              "fields are corrected. validate-models.py and validate-form-controls.py report "
              "0 problems. Evidence: defect-log.md group D"),
}

SPRINT3 = {
    "S3-24": PROJECT["PB-45"],
    "S3-25": PROJECT["PB-46"],
    "S3-23": PROJECT["PB-47"],
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true")
    args = ap.parse_args()

    wb = load_workbook(BOOK)
    edits = []

    ws = wb["Project Backlog"]
    for r in range(2, ws.max_row + 1):
        rid = str(ws.cell(row=r, column=1).value)
        if rid in PROJECT:
            ws.cell(row=r, column=13, value=DONE)
            ws.cell(row=r, column=14, value=PROJECT[rid])
            edits.append(f"Project Backlog {rid}")

    ws = wb["Sprint 3 Backlog"]
    for r in range(2, ws.max_row + 1):
        rid = str(ws.cell(row=r, column=1).value)
        if rid in SPRINT3:
            ws.cell(row=r, column=10, value=DONE)
            ws.cell(row=r, column=11, value=SPRINT3[rid])
            edits.append(f"Sprint 3 Backlog {rid}")

    for e in edits:
        print(f"  {e} -> Done")
    if args.check:
        print(f"\n{len(edits)} edit(s)  (check mode, nothing written)")
        return 0
    wb.save(BOOK)
    print(f"\n{len(edits)} edit(s) written to {BOOK.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
