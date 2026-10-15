#!/usr/bin/env python3
"""Update the SD dependency count from 37 to 40 wherever the documents quote it.

Adding the three dependencies the case study requires (see
apply-istar-consistency-fix.py) changes the headline count, and the portfolio text, the
agile record, the work-division note and the combined portfolio all state the old one.

Replacement is done at run level in the DOCX files so that bold/italic/font formatting on
the surrounding text survives - setting paragraph.text would flatten it.

usage: python update-sd-count.py [--check]
"""

import argparse
import re
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]

DOCX = [
    ROOT / "07_Portfolio_Documents" / "UFCEP4-0-3_01_Task01_02_Strategic_BPMN_and_iStar.docx",
    ROOT / "07_Portfolio_Documents" / "UFCEP4-0-3_04_Agile_Delivery_Record.docx",
    ROOT / "07_Portfolio_Documents" / "UFCEP4-0-3_Portfolio.docx",
    ROOT / "01_Agile_Planning" / "UFCEP4-0-3_Work_Division.docx",
]

MARKDOWN = [
    ROOT / "01_Agile_Planning" / "work-division.md",
]

REPLACEMENTS = [
    ("37 typed dependencies", "40 typed dependencies"),
    ("37 dependencies", "40 dependencies"),
    ("Thirty-seven dependencies", "Forty dependencies"),
]


def patch_runs(container, tally):
    for p in container.paragraphs:
        for run in p.runs:
            for old, new in REPLACEMENTS:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    tally[old] = tally.get(old, 0) + 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true")
    args = ap.parse_args()

    total = 0
    for path in DOCX:
        if not path.exists():
            print(f"  missing: {path.name}")
            continue
        doc = Document(str(path))
        tally = {}
        patch_runs(doc, tally)
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    patch_runs(cell, tally)
        total += sum(tally.values())
        print(f"  {path.name:<58} {tally if tally else 'no change'}")
        if tally and not args.check:
            doc.save(str(path))

    for path in MARKDOWN:
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8")
        before = text
        for old, new in REPLACEMENTS:
            text = text.replace(old, new)
        n = sum(before.count(o) for o, _ in REPLACEMENTS)
        total += n
        print(f"  {path.name:<58} {n} replacement(s)")
        if n and not args.check:
            path.write_text(text, encoding="utf-8")

    print(f"\n{total} replacement(s)" + ("  (check mode - nothing written)" if args.check else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main())
